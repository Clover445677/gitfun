import io
import json
import os
import re
import socket
import sys
from contextlib import asynccontextmanager
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

from dotenv import load_dotenv
from docx import Document
from fastapi import Body, Depends, FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fpdf import FPDF
from fpdf.enums import WrapMode
from openai import APITimeoutError, AsyncOpenAI, OpenAIError
from sqlalchemy.orm import Session

from models import Generation, PostPlan, PromptSettings, Version, create_database
from prompts import (
    analyze_reference_prompt,
    build_correction_prompt,
    build_copywriter_regeneration_prompt,
    build_designer_regeneration_prompt,
    build_prompt,
    default_prompt_settings,
    filter_enabled_fields,
    get_prompt_settings,
)
from schemas import (
    BriefsPayload,
    GenerationRequest,
    GenerationResponse,
    HistoryItem,
    PostPlanResponse,
    PromptSettingsPayload,
    ReferenceRequest,
    ReferenceResponse,
    RegenerationRequest,
    SheetConnectRequest,
    SubtletiesUpdate,
    VersionList,
    VersionResponse,
)

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./data/smm_briefs.db")
OLLAMA_BASE_URL = "http://localhost:11434/v1"
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "qwen3:14b")
if DATABASE_URL.startswith("sqlite:///./"):
    (BASE_DIR / "data").mkdir(exist_ok=True)
    DATABASE_URL = f"sqlite:///{BASE_DIR / DATABASE_URL.removeprefix('sqlite:///./')}"
SessionLocal = create_database(DATABASE_URL)


def log(message: str) -> None:
    """Логирование с принудительным сбросом буфера."""
    sys.stdout.write(f"{message}\n")
    sys.stdout.flush()


@asynccontextmanager
async def lifespan(_: FastAPI):
    yield


app = FastAPI(title="SMM Brief Generator", version="1.0.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def as_response(item: Generation) -> GenerationResponse:
    settings = get_prompt_settings()
    result = filter_enabled_fields(item.result, settings["field_visibility"])
    versions = item_versions(item)
    return GenerationResponse(
        id=item.id,
        created_at=item.created_at,
        input_data=GenerationRequest.model_validate(item.input_data),
        result=BriefsPayload.model_validate(result),
        current_version_id=item.current_version_id,
        versions=[
            VersionResponse(
                id=version.id, generation_id=version.generation_id, version_number=version.version_number,
                copywriter_brief=filter_enabled_fields({"copywriter_brief": version.copywriter_brief}, settings["field_visibility"])["copywriter_brief"],
                designer_brief=filter_enabled_fields({"designer_brief": version.designer_brief}, settings["field_visibility"])["designer_brief"],
                is_active=version.is_active, created_at=version.created_at, user_prompt=version.user_prompt,
                parent_version_id=version.parent_version_id,
            )
            for version in versions
        ],
    )


def item_versions(item: Generation) -> list[Version]:
    """Возвращает версии генерации в порядке создания."""
    return Session.object_session(item).query(Version).filter_by(generation_id=item.id).order_by(Version.version_number).all()


def get_active_version(item: Generation, db: Session) -> Version | None:
    if item.current_version_id:
        version = db.get(Version, item.current_version_id)
        if version and version.generation_id == item.id:
            return version
    return db.query(Version).filter_by(generation_id=item.id, is_active=True).order_by(Version.version_number.desc()).first()


def create_version(
    item: Generation,
    db: Session,
    copywriter_brief: dict,
    designer_brief: dict,
    *,
    user_prompt: str = "",
    parent_version_id: int | None = None,
) -> Version:
    """Сохраняет новую активную версию и синхронизирует устаревшее поле result."""
    db.query(Version).filter_by(generation_id=item.id, is_active=True).update({Version.is_active: False})
    last_version = db.query(Version.version_number).filter_by(generation_id=item.id).order_by(Version.version_number.desc()).first()
    version = Version(
        generation_id=item.id,
        version_number=(last_version[0] if last_version else 0) + 1,
        copywriter_brief=copywriter_brief,
        designer_brief=designer_brief,
        is_active=True,
        user_prompt=user_prompt,
        parent_version_id=parent_version_id,
    )
    result = dict(item.result)
    result["copywriter_brief"] = copywriter_brief
    result["designer_brief"] = designer_brief
    item.result = result
    db.add(version)
    db.flush()
    item.current_version_id = version.id
    return version


async def generate_briefs(request: GenerationRequest) -> dict[str, Any]:
    """Генерация ТЗ через Ollama с подробным логированием."""
    log(f"📤 Получен запрос на генерацию: {request.theme}")
    log(f"📤 Тема: {request.theme}")
    log(f"📤 Аудитория: {request.audience}")
    log(f"📤 Пол: {request.gender}")
    log(f"📤 Возраст: {request.age_from} - {request.age_to}")
    log(f"📤 Боль: {request.pain}")
    log(f"📤 Платформа: {request.platform}")
    log(f"📤 Формат: {request.content_format}")
    log(f"📤 Тон: {request.tone}")
    log(f"📤 Доп. пожелания: {request.additional_wishes}")
    log(f"📤 ai_prompt: {request.ai_prompt}")
    
    client = AsyncOpenAI(api_key="ollama", base_url=OLLAMA_BASE_URL, timeout=180.0)
    
    try:
        result = await generate_payload(client, build_prompt(request))
        if request.ai_prompt:
            log("📤 Запуск второго прохода с инструкцией пользователя...")
            result = await generate_payload(client, build_correction_prompt(result, request, request.ai_prompt))
        log("✅ JSON успешно распарсен и валидирован")
        
        return result
        
    except APITimeoutError as error:
        log("❌ Превышено время ожидания ответа от модели")
        raise HTTPException(status_code=504, detail="Превышено время ожидания ответа от модели") from error
    except (OpenAIError, ValueError, json.JSONDecodeError) as error:
        log(f"❌ Ошибка генерации: {error}")
        raise HTTPException(status_code=502, detail=f"Ошибка генерации Ollama: {error}") from error


async def generate_json(client: AsyncOpenAI, prompt_text: str) -> dict[str, Any]:
    """Выполняет один запрос к локальной Ollama и возвращает JSON-объект."""
    log(f"📤 Отправка запроса к модели {OPENAI_MODEL} (промпт: {len(prompt_text)} символов)...")
    completion = await client.chat.completions.create(
        model=OPENAI_MODEL,
        messages=[
            {"role": "system", "content": get_prompt_settings()["system_prompt"] + "\nОтвечай только валидным JSON."},
            {"role": "user", "content": prompt_text},
        ],
        temperature=0.8,
        top_p=0.95,
        extra_body={"think": False},
    )
    content = clean_json_response(extract_completion_content(completion))
    if content is None:
        raise ValueError("модель вернула невалидный JSON. Попробуйте ещё раз.")
    return fix_keys(json.loads(content))


async def generate_payload(client: AsyncOpenAI, prompt_text: str) -> dict[str, Any]:
    result = BriefsPayload.model_validate(await generate_json(client, prompt_text)).model_dump(exclude_none=True)
    settings = get_prompt_settings()
    if settings["hide_subtleties"]:
        result.pop("subtleties", None)
    return filter_enabled_fields(result, settings["field_visibility"])


def fix_keys(payload: dict[str, Any]) -> dict[str, Any]:
    """Нормализует сокращённые ключи, которые иногда возвращает модель."""
    aliases = {
        "copywriter": "copywriter_brief",
        "designer": "designer_brief",
    }
    normalized = dict(payload)
    for source_key, target_key in aliases.items():
        if source_key in normalized:
            normalized.setdefault(target_key, normalized[source_key])
            del normalized[source_key]
    return normalized


async def regenerate_brief(copywriter_brief: dict, designer_brief: dict, user_prompt: str, section: str) -> dict:
    client = AsyncOpenAI(api_key="ollama", base_url=OLLAMA_BASE_URL, timeout=180.0)
    prompt = (
        build_copywriter_regeneration_prompt(copywriter_brief, designer_brief, user_prompt)
        if section == "copywriter"
        else build_designer_regeneration_prompt(copywriter_brief, designer_brief, user_prompt)
    )
    result = await generate_json(client, prompt)
    return extract_regenerated_brief(result, section)


def extract_regenerated_brief(result: dict[str, Any], section: str) -> dict[str, Any]:
    """Извлекает перегенерированный блок из обёрнутого или прямого ответа модели."""
    key = f"{section}_brief"
    brief = result.get(key)
    if isinstance(brief, dict):
        return brief

    required_fields = (
        {"concept", "goal", "structure"}
        if section == "copywriter"
        else {"visual_task", "format", "style"}
    )
    if required_fields.issubset(result):
        log(f"⚠️ Модель вернула {key} без внешнего ключа; использован прямой объект.")
        return result

    raise HTTPException(status_code=502, detail=f"Ошибка генерации Ollama: не найден ключ {key}")


def extract_completion_content(completion: Any) -> str:
    """Извлечение контента из ответа модели (поддержка streaming и обычного режима)."""
    if not isinstance(completion, str):
        return completion.choices[0].message.content or ""

    content_parts: list[str] = []
    for line in completion.splitlines():
        if not line.startswith("data:"):
            continue
        payload = line.removeprefix("data:").strip()
        if not payload or payload == "[DONE]":
            continue
        data = json.loads(payload)
        if "error" in data:
            raise ValueError(str(data["error"]))
        for choice in data.get("choices", []):
            message = choice.get("message", {})
            delta = choice.get("delta", {})
            text = message.get("content") or delta.get("content")
            if text:
                content_parts.append(text)
    return "".join(content_parts)


def clean_json_response(raw_text: str) -> str | None:
    """Извлекает JSON-объект из ответа модели и устраняет висячие запятые."""
    start = raw_text.find("{")
    end = raw_text.rfind("}")
    if start == -1 or end == -1 or start >= end:
        return None

    candidate = raw_text[start : end + 1]
    candidate = re.sub(r",\s*(?=[}\]])", "", candidate)
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        return None
    if not isinstance(parsed, dict):
        return None
    return json.dumps(parsed, ensure_ascii=False)


def prompt_settings_response() -> dict[str, Any]:
    return get_prompt_settings()


def fetch_reference_content(url: str) -> str:
    """Загружает доступный публичный текст страницы референса."""
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise HTTPException(status_code=400, detail="Укажите корректную публичную ссылку http или https.")
    if parsed.hostname.casefold() == "t.me":
        url = parsed._replace(netloc="telegram.me").geturl()
        parsed = urlparse(url)
    try:
        import ipaddress
        from urllib.request import getproxies

        addresses = socket.getaddrinfo(parsed.hostname, None)
        if any(ipaddress.ip_address(address[4][0]).is_private or ipaddress.ip_address(address[4][0]).is_loopback for address in addresses):
            raise HTTPException(status_code=400, detail="Ссылки на внутренние адреса недоступны.")
        import requests
        from bs4 import BeautifulSoup

        proxies = getproxies()
        log(f"🔍 Используются прокси: {proxies}")
        response = requests.get(url, proxies=proxies, timeout=60, headers={"User-Agent": "Mozilla/5.0 (compatible; SMMBriefs/1.0)"})
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        content = " ".join(soup.stripped_strings)
        if len(content) < 30:
            raise ValueError("на странице недостаточно текста")
        return content[:12000]
    except HTTPException:
        raise
    except Exception as error:
        log(f"⚠️ Не удалось загрузить референс: {error}")
        raise HTTPException(status_code=422, detail="Ссылка недоступна. Проверьте URL или попробуйте позже.") from error


async def analyze_reference(content: str) -> dict[str, Any]:
    """Анализирует текст публичного референса локальной моделью."""
    client = AsyncOpenAI(api_key="ollama", base_url=OLLAMA_BASE_URL, timeout=90.0)
    try:
        return ReferenceResponse.model_validate(await generate_json(client, analyze_reference_prompt(content))).model_dump()
    except (OpenAIError, ValueError, json.JSONDecodeError) as error:
        raise HTTPException(status_code=502, detail=f"Не удалось проанализировать референс: {error}") from error


@app.get("/", include_in_schema=False)
def index():
    return FileResponse(BASE_DIR / "static" / "index.html")


@app.get("/api/prompts")
def get_prompts():
    """Возвращает сохранённые либо стандартные настройки промптов."""
    return prompt_settings_response()


@app.post("/api/prompts")
def save_prompts(payload: PromptSettingsPayload, db: Session = Depends(get_db)):
    """Сохраняет настройки промптов."""
    values = payload.model_dump()
    if values["hide_subtleties"]:
        values["example_response"].pop("subtleties", None)
    values["field_visibility"] = get_prompt_settings()["field_visibility"] | values["field_visibility"]
    setting = db.query(PromptSettings).filter_by(name="prompt_settings").first()
    if setting is None:
        db.add(PromptSettings(name="prompt_settings", content=values))
    else:
        setting.content = values
    db.commit()
    return prompt_settings_response()


@app.get("/api/prompts/example", response_model=BriefsPayload)
def get_prompt_example():
    """Возвращает актуальный валидный пример ТЗ."""
    return BriefsPayload.model_validate(get_prompt_settings()["example_response"])


@app.post("/api/prompts/example", response_model=BriefsPayload)
def save_prompt_example(example: BriefsPayload = Body(...), db: Session = Depends(get_db)):
    """Сохраняет пример ТЗ, используемый при генерации."""
    setting = db.query(PromptSettings).filter_by(name="prompt_settings").first()
    content = example.model_dump(exclude_none=True)
    if setting is None:
        values = default_prompt_settings()
        values["example_response"] = content
        db.add(PromptSettings(name="prompt_settings", content=values))
    else:
        setting.content = {**get_prompt_settings(), "example_response": content}
    db.commit()
    return example


@app.get("/api/prompts/default")
def reset_prompts(db: Session = Depends(get_db)):
    """Удаляет пользовательские переопределения и возвращает стандарты."""
    db.query(PromptSettings).filter_by(name="prompt_settings").delete()
    db.commit()
    return prompt_settings_response()


@app.post("/api/analyze-reference", response_model=ReferenceResponse)
async def analyze_reference_endpoint(request: ReferenceRequest):
    return await analyze_reference(fetch_reference_content(request.url))


async def generate_with_progress(request: GenerationRequest, db: Session):
    """Генерирует ТЗ и передаёт браузеру статус через SSE-события."""
    yield "event: progress\ndata: Анализируем тему и данные...\n\n"
    try:
        yield "event: progress\ndata: Формируем задание для копирайтера и дизайнера...\n\n"
        result = await generate_briefs(request)
        yield "event: progress\ndata: Сохраняем ТЗ в историю...\n\n"
        log("💾 Сохранение в БД...")
        item = Generation(theme=request.theme, platform=request.platform, input_data=request.model_dump(), result=result)
        db.add(item)
        db.flush()
        create_version(item, db, result["copywriter_brief"], result["designer_brief"], user_prompt=request.ai_prompt)
        db.commit()
        db.refresh(item)
        log(f"✅ Запись сохранена, ID: {item.id}")
        response = as_response(item)
        yield "event: progress\ndata: Готово!\n\n"
        yield f"event: result\ndata: {response.model_dump_json()}\n\n"
    except Exception as error:
        log(f"❌ Ошибка генерации: {error}")
        detail = error.detail if isinstance(error, HTTPException) else "Внутренняя ошибка сервера."
        yield f"event: error\ndata: {json.dumps({'detail': detail}, ensure_ascii=False)}\n\n"


@app.post("/api/generations", status_code=201)
async def create_generation(request: GenerationRequest, db: Session = Depends(get_db)):
    """Создание ТЗ с потоковой передачей статуса выполнения."""
    return StreamingResponse(
        generate_with_progress(request, db),
        media_type="text/event-stream",
        status_code=201,
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/api/generations", response_model=list[HistoryItem])
def list_generations(db: Session = Depends(get_db)):
    """Получение списка всех генераций."""
    items = db.query(Generation).order_by(Generation.created_at.desc()).all()
    return [HistoryItem(id=item.id, created_at=item.created_at, theme=item.theme, platform=item.platform) for item in items]


def get_generation_or_404(generation_id: int, db: Session) -> Generation:
    """Получение генерации по ID или вызов 404."""
    item = db.get(Generation, generation_id)
    if not item:
        raise HTTPException(status_code=404, detail="Генерация не найдена")
    return item


def parse_sheet_date(value: Any) -> datetime:
    """Преобразует дату из Google Sheets в дату публикации."""
    if isinstance(value, datetime):
        return value
    text_value = str(value or "").strip()
    for date_format in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d.%m"):
        try:
            parsed = datetime.strptime(text_value, date_format)
            if date_format == "%d.%m":
                parsed = parsed.replace(year=datetime.now().year)
            return parsed
        except ValueError:
            continue
    raise ValueError(f"не удалось распознать дату «{text_value}»")


def sheet_value(row: dict[str, Any], *names: str) -> str:
    normalized = {normalize_sheet_text(key): value for key, value in row.items()}
    for name in names:
        value = normalized.get(normalize_sheet_text(name))
        if value is not None:
            return str(value).strip()
    return ""


def normalize_sheet_text(value: Any) -> str:
    """Нормализует пробелы и регистр в ячейках Google Sheets."""
    return " ".join(str(value).replace("\u00a0", " ").split()).casefold()


def get_google_spreadsheet(sheet_url: str):
    """Открывает Google Sheets через сервисный аккаунт."""
    try:
        import gspread
        from google.oauth2.service_account import Credentials
    except ImportError as error:
        raise HTTPException(status_code=503, detail="Не установлены пакеты Google Sheets. Выполните pip install -r requirements.txt") from error

    credentials_path = Path(os.getenv("GOOGLE_APPLICATION_CREDENTIALS", BASE_DIR / "credentials.json"))
    if not credentials_path.is_file():
        raise HTTPException(status_code=503, detail="Не найден файл учётных данных Google: credentials.json")
    try:
        scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
        credentials = Credentials.from_service_account_file(credentials_path, scopes=scopes)
        return gspread.authorize(credentials).open_by_url(sheet_url)
    except Exception as error:
        raise HTTPException(status_code=502, detail=f"Не удалось открыть Google Sheets: {error}") from error


def get_sheet_names(sheet_url: str) -> list[str]:
    """Возвращает названия всех вкладок доступной Google-таблицы."""
    return [worksheet.title for worksheet in get_google_spreadsheet(sheet_url).worksheets()]


def get_sheet_rows(sheet_url: str, sheet_name: str = "Сентябрь") -> list[list[str]]:
    """Получает строки указанного листа Google Sheets через сервисный аккаунт."""
    try:
        import gspread
    except ImportError as error:
        raise HTTPException(status_code=503, detail="Не установлены пакеты Google Sheets. Выполните pip install -r requirements.txt") from error
    try:
        spreadsheet = get_google_spreadsheet(sheet_url)
        try:
            worksheet = spreadsheet.worksheet(sheet_name)
        except gspread.WorksheetNotFound as error:
            gid = parse_qs(urlparse(sheet_url).query).get("gid", [None])[0]
            worksheet = next((item for item in spreadsheet.worksheets() if str(item.id) == gid), None)
            if worksheet is None:
                available_names = ", ".join(item.title for item in spreadsheet.worksheets())
                raise HTTPException(
                    status_code=404,
                    detail=f"Лист «{sheet_name}» не найден. Доступные листы: {available_names or 'нет'}",
                ) from error
        return worksheet.get_all_values()
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=502, detail=f"Не удалось прочитать Google Sheets: {error}") from error


def import_plan_rows(sheet_url: str, rows: list[list[str]], db: Session, sheet_name: str = "Сентябрь") -> int:
    """Находит заголовки в одной или нескольких строках и импортирует контент-план."""
    required_headers = {normalize_sheet_text(name) for name in ("Дата", "Тема поста", "Описание поста")}
    known_headers = {normalize_sheet_text(name) for name in ("Дата", "Тема поста", "Тема", "Рубрика", "Описание поста", "Описание", "Сторис", "Идея", "Идея/Доп. поле", "Статус", "Комментарий клиента", "Комментарий")}
    header_index: int | None = None
    headers: list[str] = []

    for start_index in range(len(rows)):
        candidate_headers: list[str] = []
        for end_index in range(start_index, min(start_index + 3, len(rows))):
            row = rows[end_index]
            if len(candidate_headers) < len(row):
                candidate_headers.extend("" for _ in range(len(row) - len(candidate_headers)))
            for column_index, value in enumerate(row):
                text_value = " ".join(str(value).replace("\u00a0", " ").split())
                if text_value:
                    candidate_headers[column_index] = (
                        text_value
                        if normalize_sheet_text(text_value) in known_headers
                        else " ".join(value for value in (candidate_headers[column_index], text_value) if value)
                    )
            if required_headers.issubset({normalize_sheet_text(header) for header in candidate_headers}):
                header_index = end_index
                headers = candidate_headers
                break
        if header_index is not None:
            break

    if header_index is None:
        log("⚠️ Не найдены заголовки «Дата», «Тема поста» и «Описание поста» в первых строках листа.")
        return 0

    imported = 0
    for row_index, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
        row = {
            header: values[column_index] if column_index < len(values) else ""
            for column_index, header in enumerate(headers)
            if header
        }
        theme = sheet_value(row, "Тема поста", "Тема")
        date_value = sheet_value(row, "Дата")
        if not theme or not date_value:
            continue
        try:
            publish_date = parse_sheet_date(date_value)
        except ValueError as error:
            log(f"⚠️ Строка {row_index} пропущена: {error}")
            continue
        post = db.query(PostPlan).filter_by(sheet_url=sheet_url, sheet_name=sheet_name, row_index=row_index).first()
        if post is None:
            post = PostPlan(sheet_url=sheet_url, sheet_name=sheet_name, row_index=row_index, publish_date=publish_date, theme=theme)
            db.add(post)
        post.publish_date = publish_date
        post.theme = theme
        post.rubric = sheet_value(row, "Рубрика")
        post.description = sheet_value(row, "Описание поста", "Описание")
        post.idea = sheet_value(row, "Сторис", "Идея/Доп. поле", "Идея") or None
        post.status = sheet_value(row, "Статус")
        post.client_comment = sheet_value(row, "Комментарий клиента", "Комментарий") or None
        imported += 1
    db.commit()
    return imported


def post_plan_response(post: PostPlan) -> PostPlanResponse:
    return PostPlanResponse.model_validate(post).model_copy(
        update={"brief_status": "created" if post.brief_id else "not_created"}
    )


def build_plan_request(post: PostPlan) -> GenerationRequest:
    return GenerationRequest(
        theme=post.theme,
        pain=post.description,
        platform="Instagram",
        content_format="Авто",
        tone="Экспертный",
        ai_prompt=(
            f"Рубрика: {post.rubric or 'Не указана'}\n"
            f"Идея: {post.idea or 'Не указана'}\n"
            f"Комментарий клиента: {post.client_comment or 'Не указан'}"
        ),
    )


async def create_plan_generation(post: PostPlan, db: Session) -> Generation:
    request = build_plan_request(post)
    result = await generate_briefs(request)
    item = Generation(theme=post.theme, platform=request.platform, input_data=request.model_dump(), result=result)
    db.add(item)
    db.flush()
    create_version(item, db, result["copywriter_brief"], result["designer_brief"], user_prompt=request.ai_prompt)
    post.brief_id = item.id
    db.commit()
    db.refresh(item)
    db.refresh(post)
    return item


@app.post("/api/plan/connect")
def connect_sheet(request: SheetConnectRequest, db: Session = Depends(get_db)):
    """Подключает Google Sheets и импортирует строки указанного листа."""
    imported = import_plan_rows(request.sheet_url, get_sheet_rows(request.sheet_url, request.sheet_name), db, request.sheet_name)
    message = "" if imported else "Не найдено строк с заполненными колонками «Дата» и «Тема поста». Проверьте выбранный лист и структуру данных."
    return {"imported": imported, "message": message}


@app.get("/api/plan/sheets")
def list_sheet_names(sheet_url: str = Query(min_length=1, max_length=500)):
    """Возвращает вкладки таблицы для выбора в интерфейсе."""
    return {"sheets": get_sheet_names(sheet_url)}


@app.get("/api/plan/posts", response_model=list[PostPlanResponse])
def get_plan_posts(
    filter: str = Query(default="all", pattern="^(all|pending|ready|published)$"),
    db: Session = Depends(get_db),
):
    """Возвращает контент-план с фильтрацией по состоянию ТЗ."""
    query = db.query(PostPlan)
    if filter == "pending":
        query = query.filter(PostPlan.brief_id.is_(None))
    elif filter == "ready":
        query = query.filter(PostPlan.brief_id.is_not(None))
    elif filter == "published":
        query = query.filter(PostPlan.status.ilike("%опублик%"))
    return [post_plan_response(post) for post in query.order_by(PostPlan.publish_date, PostPlan.id).all()]


@app.post("/api/plan/generate/{post_id}")
async def generate_plan_brief(post_id: int, db: Session = Depends(get_db)):
    """Создаёт ТЗ для одного поста контент-плана."""
    post = db.get(PostPlan, post_id)
    if not post:
        raise HTTPException(status_code=404, detail="Пост не найден")
    item = await create_plan_generation(post, db)
    return {"id": item.id, "brief_id": item.id}


@app.post("/api/plan/regenerate/{post_id}")
async def regenerate_plan_brief_from_scratch(post_id: int, db: Session = Depends(get_db)):
    """Удаляет старое ТЗ поста вместе с версиями и создаёт независимую версию 1."""
    post = db.get(PostPlan, post_id)
    if not post:
        raise HTTPException(status_code=404, detail="Пост не найден")
    if post.brief_id:
        old_generation_id = post.brief_id
        post.brief_id = None
        db.query(Version).filter_by(generation_id=old_generation_id).delete()
        old_item = db.get(Generation, old_generation_id)
        if old_item:
            db.delete(old_item)
        db.commit()
    item = await create_plan_generation(post, db)
    return {"id": item.id, "brief_id": item.id, "message": "ТЗ создано с нуля как версия 1."}


async def generate_posts(posts: list[PostPlan], db: Session) -> dict[str, Any]:
    created: list[dict[str, int]] = []
    errors: list[dict[str, Any]] = []
    for post in posts:
        try:
            item = await create_plan_generation(post, db)
            created.append({"post_id": post.id, "brief_id": item.id})
        except Exception as error:
            db.rollback()
            log(f"❌ Не удалось создать ТЗ для поста {post.id}: {error}")
            errors.append({"post_id": post.id, "error": str(error)})
    return {"created": len(created), "posts": created, "errors": errors}


@app.post("/api/plan/generate-all")
async def generate_all_briefs(db: Session = Depends(get_db)):
    """Создаёт ТЗ для всех постов, у которых его ещё нет."""
    posts = db.query(PostPlan).filter(PostPlan.brief_id.is_(None)).order_by(PostPlan.publish_date).all()
    return await generate_posts(posts, db)


@app.post("/api/plan/generate-month")
async def generate_month_briefs(db: Session = Depends(get_db)):
    """Создаёт ТЗ для постов на ближайшие 30 дней."""
    today = datetime.now().date()
    start = datetime.combine(today, time.min)
    end = datetime.combine(today + timedelta(days=30), time.max)
    posts = db.query(PostPlan).filter(PostPlan.brief_id.is_(None), PostPlan.publish_date.between(start, end)).order_by(PostPlan.publish_date).all()
    return await generate_posts(posts, db)


@app.post("/api/plan/sync")
def sync_plan(db: Session = Depends(get_db)):
    """Синхронизирует все ранее подключённые листы Google Sheets."""
    sheets = db.query(PostPlan.sheet_url, PostPlan.sheet_name).distinct().all()
    if not sheets:
        raise HTTPException(status_code=400, detail="Сначала подключите Google Sheets")
    imported = sum(
        import_plan_rows(sheet_url, get_sheet_rows(sheet_url, sheet_name), db, sheet_name)
        for sheet_url, sheet_name in sheets
    )
    return {"imported": imported}


@app.get("/api/generations/{generation_id}/versions", response_model=VersionList)
def list_versions(generation_id: int, db: Session = Depends(get_db)):
    item = get_generation_or_404(generation_id, db)
    settings = get_prompt_settings()
    versions = []
    for version in item_versions(item):
        briefs = filter_enabled_fields(
            {"copywriter_brief": version.copywriter_brief, "designer_brief": version.designer_brief},
            settings["field_visibility"],
        )
        versions.append(VersionResponse(
            id=version.id, generation_id=version.generation_id, version_number=version.version_number,
            copywriter_brief=briefs["copywriter_brief"], designer_brief=briefs["designer_brief"],
            is_active=version.is_active, created_at=version.created_at, user_prompt=version.user_prompt,
            parent_version_id=version.parent_version_id,
        ))
    return VersionList(versions=versions)


@app.get("/api/generations/{generation_id}/versions/{version_id}", response_model=VersionResponse)
def get_version(generation_id: int, version_id: int, db: Session = Depends(get_db)):
    version = db.get(Version, version_id)
    if not version or version.generation_id != generation_id:
        raise HTTPException(status_code=404, detail="Версия не найдена")
    return VersionResponse.model_validate(version)


async def regenerate_generation_section(generation_id: int, request: RegenerationRequest, section: str, db: Session) -> VersionResponse:
    item = get_generation_or_404(generation_id, db)
    active = get_active_version(item, db)
    copywriter_brief = active.copywriter_brief if active else item.result["copywriter_brief"]
    designer_brief = active.designer_brief if active else item.result["designer_brief"]
    updated_brief = await regenerate_brief(copywriter_brief, designer_brief, request.prompt, section)
    filtered = filter_enabled_fields({f"{section}_brief": updated_brief}, get_prompt_settings()["field_visibility"])
    updated_brief = filtered[f"{section}_brief"]
    if section == "copywriter":
        copywriter_brief = updated_brief
    else:
        designer_brief = updated_brief
    version = create_version(
        item,
        db,
        copywriter_brief,
        designer_brief,
        user_prompt=request.prompt, parent_version_id=active.id if active else None,
    )
    db.commit()
    db.refresh(version)
    return VersionResponse.model_validate(version)


@app.post("/api/generations/{generation_id}/regenerate/copywriter", response_model=VersionResponse)
async def regenerate_copywriter(generation_id: int, request: RegenerationRequest, db: Session = Depends(get_db)):
    return await regenerate_generation_section(generation_id, request, "copywriter", db)


@app.post("/api/generations/{generation_id}/regenerate/designer", response_model=VersionResponse)
async def regenerate_designer(generation_id: int, request: RegenerationRequest, db: Session = Depends(get_db)):
    return await regenerate_generation_section(generation_id, request, "designer", db)


@app.get("/api/generations/{generation_id}", response_model=GenerationResponse)
def get_generation(generation_id: int, db: Session = Depends(get_db)):
    """Получение конкретной генерации по ID."""
    return as_response(get_generation_or_404(generation_id, db))


@app.delete("/api/generations/{generation_id}", status_code=204)
def delete_generation(generation_id: int, db: Session = Depends(get_db)):
    """Удаление генерации по ID."""
    item = get_generation_or_404(generation_id, db)
    db.query(Version).filter_by(generation_id=item.id).delete()
    db.delete(item)
    db.commit()


@app.put("/api/generations/{generation_id}/subtleties", response_model=GenerationResponse)
def update_subtleties(generation_id: int, update: SubtletiesUpdate, db: Session = Depends(get_db)):
    """Обновление тонкостей в существующей генерации."""
    item = get_generation_or_404(generation_id, db)
    result = dict(item.result)
    if get_prompt_settings()["hide_subtleties"] or "subtleties" not in result:
        raise HTTPException(status_code=409, detail="Блок тонкостей отключён в настройках промпта")
    subtleties = dict(result["subtleties"])
    subtleties.update(update.model_dump())
    result["subtleties"] = subtleties
    item.result = result
    db.commit()
    db.refresh(item)
    return as_response(item)


# Метки для экспорта
SECTION_LABELS = {
    "key_message": "Единый ключевой месседж", "deadline": "Срок сдачи", "revisions": "Количество правок",
    "delivery_format": "Формат сдачи", "priority": "Приоритет", "concept": "Общая концепция", "goal": "Цель поста",
    "tone": "Тон общения", "volume": "Объем", "keywords": "Ключевые слова/фразы", "stop_words": "Стоп-слова",
    "recommendations": "Дополнительные рекомендации", "visual_task": "Общая визуальная задача", "format": "Формат",
    "palettes": "Цветовая гамма", "required_elements": "Обязательные элементы", "composition": "Композиция",
    "style": "Атмосфера/стиль", "references": "Референсы по визуалу", "technical_requirements": "Технические требования",
    "headlines": "Заголовок (H1), 3 варианта", "hooks": "Вступление (крючок), 2 варианта",
    "main_points": "Основная часть", "conclusion": "Вывод/резюме", "ctas": "Призыв к действию (CTA), 2 варианта",
}


def add_mapping_to_doc(document: Document, data: dict, level: int = 2) -> None:
    """Рекурсивное добавление данных в документ DOCX."""
    for key, value in data.items():
        label = SECTION_LABELS.get(key, key)
        if isinstance(value, dict):
            document.add_heading(label, level=level)
            add_mapping_to_doc(document, value, level + 1)
        elif isinstance(value, list):
            document.add_heading(label, level=level)
            for line in value:
                document.add_paragraph(str(line), style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(str(value))


def plain_text(data: dict) -> str:
    """Преобразование данных в текстовый формат для PDF."""
    lines: list[str] = []
    for key, value in data.items():
        label = SECTION_LABELS.get(key, key)
        if isinstance(value, dict):
            lines.extend([label.upper(), plain_text(value)])
        elif isinstance(value, list):
            lines.append(f"{label}:\n" + "\n".join(f"• {item}" for item in value))
        else:
            lines.append(f"{label}: {value}")
    return "\n\n".join(lines)


def export_text(item: Generation, section: str | None = None) -> str:
    """Экспорт данных генерации в текстовый формат."""
    result = filter_enabled_fields(item.result, get_prompt_settings()["field_visibility"])
    parts = [f"ТЗ ДЛЯ SMM: {item.theme}"]
    if not get_prompt_settings()["hide_subtleties"] and result.get("subtleties"):
        parts.append("ТОНКОСТИ\n" + plain_text(result["subtleties"]))
    if section in (None, "copywriter"):
        parts.append("ТЗ КОПИРАЙТЕРУ\n" + plain_text(result["copywriter_brief"]))
    if section in (None, "designer"):
        parts.append("ТЗ ДИЗАЙНЕРУ\n" + plain_text(result["designer_brief"]))
    return "\n\n".join(parts)


def export_docx_bytes(item: Generation, section: str | None = None) -> bytes:
    """Экспорт данных генерации в формат DOCX."""
    result = filter_enabled_fields(item.result, get_prompt_settings()["field_visibility"])
    document = Document()
    document.add_heading(f"ТЗ для SMM: {item.theme}", 0)
    if not get_prompt_settings()["hide_subtleties"] and result.get("subtleties"):
        document.add_heading("Тонкости", 1)
        add_mapping_to_doc(document, result["subtleties"])
    if section in (None, "copywriter"):
        document.add_heading("ТЗ копирайтеру", 1)
        add_mapping_to_doc(document, result["copywriter_brief"])
    if section in (None, "designer"):
        document.add_heading("ТЗ дизайнеру", 1)
        add_mapping_to_doc(document, result["designer_brief"])
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def export_pdf_bytes(item: Generation, section: str | None = None) -> bytes:
    """Экспорт данных генерации в формат PDF."""
    font_candidates = [os.getenv("FONT_PATH"), "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    font_path = next((candidate for candidate in font_candidates if candidate and Path(candidate).is_file()), None)
    if not font_path:
        raise HTTPException(status_code=500, detail="Не найден шрифт с поддержкой кириллицы. Укажите FONT_PATH в .env")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("BriefFont", "", font_path)
    pdf.set_font("BriefFont", size=11)
    for line in export_text(item, section).splitlines():
        pdf.multi_cell(0, 7, line or " ", new_x="LMARGIN", new_y="NEXT", wrapmode=WrapMode.CHAR)
    return bytes(pdf.output())


def download_response(content: bytes, media_type: str, filename: str) -> StreamingResponse:
    """Формирование ответа для скачивания файла."""
    return StreamingResponse(io.BytesIO(content), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@app.get("/api/export/copywriter/docx/{generation_id}")
def export_copywriter_docx(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт ТЗ копирайтера в DOCX."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_docx_bytes(item, "copywriter"), DOCX_MEDIA_TYPE, f"smm-copywriter-{item.id}.docx")


@app.get("/api/export/copywriter/pdf/{generation_id}")
def export_copywriter_pdf(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт ТЗ копирайтера в PDF."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_pdf_bytes(item, "copywriter"), "application/pdf", f"smm-copywriter-{item.id}.pdf")


@app.get("/api/export/designer/docx/{generation_id}")
def export_designer_docx(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт ТЗ дизайнера в DOCX."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_docx_bytes(item, "designer"), DOCX_MEDIA_TYPE, f"smm-designer-{item.id}.docx")


@app.get("/api/export/designer/pdf/{generation_id}")
def export_designer_pdf(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт ТЗ дизайнера в PDF."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_pdf_bytes(item, "designer"), "application/pdf", f"smm-designer-{item.id}.pdf")


@app.get("/api/generations/{generation_id}/export/docx")
def export_docx(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт полного ТЗ в DOCX."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_docx_bytes(item), DOCX_MEDIA_TYPE, f"smm-brief-{item.id}.docx")


@app.get("/api/generations/{generation_id}/export/pdf")
def export_pdf(generation_id: int, db: Session = Depends(get_db)):
    """Экспорт полного ТЗ в PDF."""
    item = get_generation_or_404(generation_id, db)
    return download_response(export_pdf_bytes(item), "application/pdf", f"smm-brief-{item.id}.pdf")
